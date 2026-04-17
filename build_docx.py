#!/usr/bin/env python3
"""Build Word documents from cleaned OCR text files."""

import os
import glob
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS = [
    {
        "stem": "consulting_1973",
        "title": "20 Hornton Street Flat 1 License for Consulting -- 1973",
        "pages": 4,
    },
    {
        "stem": "sublease_1973",
        "title": "20 Hornton Street Flat 1 Sublease to Newnhams -- 1973",
        "pages": 26,
    },
    {
        "stem": "conversion_1972",
        "title": "20 Hornton Street License for Conversion -- 1972",
        "pages": 11,
    },
]

def build_docx(doc_info, cleaned_dir="work/cleaned", output_dir="output"):
    """Build a single Word document from cleaned OCR text files."""
    stem = doc_info["stem"]
    title = doc_info["title"]
    num_pages = doc_info["pages"]

    os.makedirs(output_dir, exist_ok=True)

    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Georgia'
    font.size = Pt(11)

    sections = doc.sections
    for section in sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for page_num in range(1, num_pages + 1):
        txt_path = os.path.join(cleaned_dir, stem, f"page_{page_num:03d}.txt")

        if not os.path.exists(txt_path):
            fallback = os.path.join("work/ocr", stem, f"page_{page_num:03d}.txt")
            if os.path.exists(fallback):
                txt_path = fallback
            else:
                continue

        with open(txt_path, 'r', encoding='utf-8') as f:
            text = f.read().strip()

        if not text:
            continue

        if page_num > 1:
            doc.add_page_break()

        for para_text in text.split('\n\n'):
            para_text = para_text.strip()
            if not para_text:
                continue

            lines = para_text.replace('\n', ' ')
            lines = ' '.join(lines.split())

            p = doc.add_paragraph(lines)
            p.style = doc.styles['Normal']

    output_path = os.path.join(output_dir, f"{title}.docx")
    doc.save(output_path)
    print(f"Saved: {output_path}")
    return output_path


if __name__ == '__main__':
    for doc_info in DOCS:
        try:
            build_docx(doc_info)
        except Exception as e:
            print(f"Error building {doc_info['stem']}: {e}")
