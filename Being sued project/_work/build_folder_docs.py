#!/usr/bin/env python3
"""Build Word documents for each folder of the Being sued project."""

import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path("/home/user/pdf-text/Being sued project")
OCR = BASE / "_work" / "ocr"
EXTRACTED = BASE / "_work" / "extracted"
OUTPUT = BASE / "_output"


def new_doc(title):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def add_file_section(doc, header, text, italian=False):
    doc.add_heading(header, level=2)
    if italian:
        p = doc.add_paragraph()
        r = p.add_run("[Italian original below]")
        r.italic = True
    for chunk in text.split("\n\n"):
        chunk = chunk.strip()
        if not chunk:
            continue
        for line in chunk.split("\n"):
            doc.add_paragraph(line)
        doc.add_paragraph("")


def read_file(path):
    with open(path, 'r', encoding='utf-8') as f:
        return f.read()


def build_advice_doc():
    doc = new_doc("Advice Folder — Lawyer Email Thread (Italian)")
    doc.add_paragraph("Email thread between Grahame McGirr (Director, Naissance UK Ltd) and Avv. Primo Belardi (current Italian lawyer) on the Gavioli bankruptcy lawsuit strategy. Screenshots from mobile Gmail client. All content in Italian.")
    doc.add_paragraph("")
    for i in range(11, 20):
        p = OCR / f"advice_shared_image_{i}.txt"
        if p.exists():
            add_file_section(doc, f"Image {i}", read_file(p), italian=True)
    out = OUTPUT / "01_Advice_lawyer_emails.docx"
    doc.save(out)
    print(f"  → {out.name}")


def build_email1_doc():
    doc = new_doc("Email 1 — Grahame's Position Letter + Fiorilli's Istanza")
    doc.add_heading("Part A: Grahame McGirr's email to Avv. Belardi", level=2)
    doc.add_paragraph("Source: 'email 1 text.docx' — full email thread between Grahame and Belardi.")
    doc.add_paragraph("")
    email_text = read_file(EXTRACTED / "email1_text.txt")
    for line in email_text.split("\n"):
        doc.add_paragraph(line)

    doc.add_page_break()
    doc.add_heading("Part B: Fiorilli's 2018 Istanza di Ammissione al Passivo (Italian, 3 pages)", level=2)
    doc.add_paragraph("Source: 'doc. 05a ISTANZA AMMISSIONE AL PASSIVO.pdf' — the 18.12.2018 filing by Avv. Paolo Fiorilli (Rome) on behalf of Naissance. Filed at Tribunale Civile di Siena, Fallimento n.35/2018.")
    doc.add_paragraph("")
    for i in range(1, 4):
        p = OCR / f"istanza_page_{i:03d}.txt"
        if p.exists():
            add_file_section(doc, f"Istanza — Page {i}", read_file(p), italian=True)

    out = OUTPUT / "02_Email1_position_and_Istanza.docx"
    doc.save(out)
    print(f"  → {out.name}")


def build_fx_receipts_doc():
    doc = new_doc("FX Receipts — Wire Transfers to Argo Ge.Re.Cre. S.r.l.")
    doc.add_paragraph("Emails from Thomas Exchange UK confirming EUR wire transfers from Grahame McGirr / Naissance to Argo Ge.Re.Cre. S.r.l. (Italy), April-May 2018.")
    doc.add_paragraph("")
    summary_path = BASE / "_work" / "translated" / "fx_receipts_summary.txt"
    if summary_path.exists():
        doc.add_heading("Summary", level=2)
        for line in read_file(summary_path).split("\n"):
            doc.add_paragraph(line)

    doc.add_page_break()
    doc.add_heading("Raw emails (extracted)", level=2)
    for eml_path in sorted(EXTRACTED.glob("fx_receipts_*.txt")):
        add_file_section(doc, eml_path.name, read_file(eml_path))

    out = OUTPUT / "03_FX_Receipts_payments.docx"
    doc.save(out)
    print(f"  → {out.name}")


def build_other_ai_doc():
    doc = new_doc("Other AI — ChatGPT/WhatsApp Strategy Discussion")
    doc.add_paragraph("Screenshots from a WhatsApp conversation where Grahame McGirr forwards/pastes AI-generated legal analysis (likely ChatGPT). All content in English.")
    doc.add_paragraph("")
    for i in range(11, 16):
        p = OCR / f"otherai_shared_image_{i}.txt"
        if p.exists():
            add_file_section(doc, f"Image {i}", read_file(p))
    out = OUTPUT / "04_Other_AI_strategy.docx"
    doc.save(out)
    print(f"  → {out.name}")


def build_ricorso_doc():
    doc = new_doc("Trustee's Lawsuit (Ricorso) — Against Naissance UK Limited")
    doc.add_paragraph("19 images of the formal complaint filed 30.12.2025 at Tribunale di Siena (Case 2505/2025 R.G.) by Avv. Fabio Finetti on behalf of Curatela Fallimento Agricola Gavioli S.r.l. Hearing: 14 May 2026. Amount claimed: €57,536.75. Documents include Italian original, English translation, judge's orders, and service-of-process records.")
    doc.add_paragraph("")
    doc.add_heading("Key facts", level=2)
    for bullet in [
        "Case no. 2505/2025 R.G. — Tribunale Ordinario di Siena, Sezione Unica Civile",
        "Judge: Dott. Marianna Serrao",
        "Plaintiff's lawyer: Avv. Fabio Finetti (Via del Paradiso 32, Siena)",
        "Curator: Dott. Stefano Scarpellini",
        "Filed: 30.12.2025 — Ricorso ex Art. 281 decies/undecies C.P.C.",
        "Hearing: originally 12.03.2026, postponed to 14.05.2026 at 10:00 (decree 07.01.2026)",
        "Amount claimed: €57,536.75 + legal interest from 01.07.2025 + costs",
        "Unified court fee: €759.00",
        "Service: Hague Convention 15.11.1965 via UK Central Authority (Royal Courts of Justice)",
        "Key case law cited: Cass. 28.09.2018 n. 23482; Cass. 20.04.2022 n. 12673",
        "Legal basis: Art. 110 Legge Fallimentare (final ranking of claims in bankruptcy)",
    ]:
        doc.add_paragraph(bullet, style='List Bullet')

    doc.add_page_break()
    doc.add_heading("Raw OCR of all 19 pages (images 11-29)", level=2)
    for i in range(11, 30):
        p = OCR / f"receipt_shared_image_{i}.txt"
        if p.exists():
            add_file_section(doc, f"Image {i}", read_file(p))

    out = OUTPUT / "05_Trustee_Ricorso_lawsuit.docx"
    doc.save(out)
    print(f"  → {out.name}")


if __name__ == '__main__':
    OUTPUT.mkdir(parents=True, exist_ok=True)
    print("Building folder Word documents...")
    build_advice_doc()
    build_email1_doc()
    build_fx_receipts_doc()
    build_other_ai_doc()
    build_ricorso_doc()
    print("Done.")
