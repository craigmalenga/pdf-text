#!/usr/bin/env python3
"""Extract text from .eml email files and .docx Word docs."""

import os
import email
from email import policy
from email.parser import BytesParser
from pathlib import Path
from docx import Document

BASE = Path("/home/user/pdf-text/Being sued project")
OUT = BASE / "_work" / "extracted"


def extract_eml(path):
    """Extract readable text from an .eml file."""
    with open(path, 'rb') as f:
        msg = BytesParser(policy=policy.default).parse(f)

    parts = []
    parts.append(f"From: {msg.get('From', '')}")
    parts.append(f"To: {msg.get('To', '')}")
    parts.append(f"Cc: {msg.get('Cc', '')}")
    parts.append(f"Date: {msg.get('Date', '')}")
    parts.append(f"Subject: {msg.get('Subject', '')}")
    parts.append("")

    body = msg.get_body(preferencelist=('plain', 'html'))
    if body:
        content = body.get_content()
        parts.append(content)
    else:
        for part in msg.walk():
            if part.get_content_type() == 'text/plain':
                parts.append(part.get_content())

    attachments = []
    for part in msg.iter_attachments():
        fname = part.get_filename()
        if fname:
            attachments.append(fname)
    if attachments:
        parts.append("")
        parts.append(f"[Attachments: {', '.join(attachments)}]")

    return '\n'.join(parts)


def extract_docx(path):
    """Extract readable text from a .docx file."""
    doc = Document(path)
    lines = []
    for para in doc.paragraphs:
        lines.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = ' | '.join(cell.text for cell in row.cells)
            lines.append(row_text)
    return '\n'.join(lines)


def main():
    os.makedirs(OUT, exist_ok=True)

    eml_dir = BASE / "FW_ 57,000 euro paid fx receipts"
    for eml_path in sorted(eml_dir.glob("*.eml")):
        print(f"Extracting {eml_path.name}")
        text = extract_eml(eml_path)
        out_name = eml_path.stem.replace(' ', '_').replace('[', '').replace(']', '') + ".txt"
        out_path = OUT / f"fx_receipts_{out_name}"
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(text)
        print(f"  → {out_path.name} ({len(text)} chars)")

    docx_path = BASE / "Email 1" / "email 1 text.docx"
    if docx_path.exists():
        print(f"Extracting {docx_path.name}")
        text = extract_docx(docx_path)
        out_path = OUT / "email1_text.txt"
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(text)
        print(f"  → {out_path.name} ({len(text)} chars)")


if __name__ == '__main__':
    main()
